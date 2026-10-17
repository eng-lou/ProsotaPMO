"""Build the Maritime user guide from its maintained Markdown source.

Run with the Codex bundled Python (python-docx). No application or project data
is read or changed. Refresh page-reference fields in Word before PDF rendering.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs' / 'MARITIME_USER_GUIDE.md'
OUTPUT = ROOT / 'docs' / 'guides' / 'Prosota_Maritime_User_Guide.docx'
BLACK = '000000'
BLUE = '1D3551'


def child(parent, tag, **attrs):
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(key), str(value))
    parent.append(el)
    return el


def set_font(style, size, bold=False, color=BLACK):
    style.font.name = 'Calibri'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    color_el = rpr.find(qn('w:color'))
    if color_el is not None:
        for a in ('w:themeColor', 'w:themeTint', 'w:themeShade'):
            color_el.attrib.pop(qn(a), None)


def bookmark(p, name, number):
    child(p._p, 'w:bookmarkStart', **{'w:id':number, 'w:name':name})
    child(p._p, 'w:bookmarkEnd', **{'w:id':number})


def hyperlink(p, label, target=None, anchor=None):
    link = OxmlElement('w:hyperlink')
    if anchor:
        link.set(qn('w:anchor'), anchor)
    else:
        link.set(qn('r:id'), p.part.relate_to(target, RT.HYPERLINK, is_external=True))
    r = child(link, 'w:r')
    props = child(r, 'w:rPr')
    child(props, 'w:color', **{'w:val':'235C72'})
    child(props, 'w:u', **{'w:val':'single'})
    child(r, 'w:t').text = label
    p._p.append(link)


def inline(p, text):
    pattern = r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))'
    for part in re.split(pattern, text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            run = p.add_run(part[1:-1]); run.font.name = 'Consolas'; run.font.size = Pt(9.5)
        elif part.startswith('['):
            m = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', part)
            if m:
                hyperlink(p, m[1], target=m[2])
            else:
                p.add_run(part)
        else:
            p.add_run(part)


def field(p, code, placeholder='1'):
    run = p.add_run()
    child(run._r, 'w:fldChar', **{'w:fldCharType':'begin'})
    text = child(run._r, 'w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = f' {code} '
    child(run._r, 'w:fldChar', **{'w:fldCharType':'separate'})
    child(run._r, 'w:t').text = placeholder
    child(run._r, 'w:fldChar', **{'w:fldCharType':'end'})


def add_table(doc, rows):
    columns = len(rows[0])
    widths = {2:[1.85,4.92], 3:[1.85,2.25,2.67]}.get(columns, [6.77/columns]*columns)
    # Keep the cost example's formula column wide enough to read naturally.
    if rows[0][0] == 'Starter cost component':
        widths = [1.8,3.65,1.32]
    table = doc.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    pr = table._tbl.tblPr
    borders = child(pr, 'w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        child(borders, f'w:{side}', **{'w:val':'single','w:sz':'5','w:color':'D9D9D9'})
    margins = child(pr, 'w:tblCellMar')
    for side, val in [('top',90),('bottom',90),('left',110),('right',110)]:
        child(margins, f'w:{side}', **{'w:w':val,'w:type':'dxa'})
    for i, values in enumerate(rows):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        child(trpr, 'w:cantSplit')
        if i == 0:
            child(trpr, 'w:tblHeader')
        for j, (cell, value) in enumerate(zip(row.cells,values)):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            child(tcpr, 'w:shd', **{'w:fill': BLUE if i == 0 else ('F1F5F8' if i % 2 == 0 else 'FFFFFF')})
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            if i == 0:
                p.paragraph_format.keep_with_next = True
            inline(p, value)
            for r in p.runs:
                r.font.size = Pt(10)
                if i == 0:
                    r.bold = True; r.font.color.rgb = RGBColor(255,255,255)
            if rows[0][0] == 'Starter cost component' and j == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)
    after.paragraph_format.line_spacing = Pt(2)
    after.paragraph_format.space_before = Pt(0)


def build():
    text = SOURCE.read_text(encoding='utf-8')
    lines = text.splitlines()
    headings = [line[3:] for line in lines if line.startswith('## ')]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27); section.page_height = Inches(11.69)
    section.top_margin = Inches(.7); section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.75); section.right_margin = Inches(.75)
    section.header_distance = Inches(.28); section.footer_distance = Inches(.3)
    section.different_first_page_header_footer = True
    for name, size, bold in [('Normal',11,False),('Title',28,True),('Subtitle',12,False),('Heading 1',18,True),('Heading 2',13,True),('Heading 3',11,True)]:
        st = doc.styles[name]; set_font(st,size,bold)
        if name.startswith('Heading'):
            st.paragraph_format.keep_with_next = True
            st.paragraph_format.space_before = Pt(13)
            st.paragraph_format.space_after = Pt(7)
    normal = doc.styles['Normal'].paragraph_format
    normal.space_after = Pt(7); normal.line_spacing = 1.09
    normal.widow_control = True
    for name in ['Title','Subtitle']:
        pr=doc.styles[name].element.find(qn('w:pPr'))
        if pr is not None:
            for b in list(pr.findall(qn('w:pBdr'))): pr.remove(b)
    header=section.header.paragraphs[0]
    header.add_run('Prosota Maritime User Guide').font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor(0,0,0)
    footer=section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Version 1.0  |  ')
    field(footer,'PAGE')
    for r in footer.runs:r.font.size=Pt(9)
    doc.core_properties.title='Prosota Maritime User Guide'
    doc.core_properties.subject='Using the persistent Maritime workspace'
    doc.core_properties.author='Prosota'
    doc.core_properties.keywords='Maritime, user guide, operations, safety, delivery'
    child(doc.settings.element, 'w:updateFields', **{'w:val':'true'})

    toc_added=False; section_count=0; i=0
    while i < len(lines):
        line=lines[i].strip()
        if not line:
            i+=1;continue
        if line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title');i+=1;continue
        if line.startswith('Version 1.0 |'):
            doc.add_paragraph(line.replace(' | ','\n'),'Subtitle');i+=1;continue
        if line.startswith('## '):
            if not toc_added:
                doc.add_page_break()
                p=doc.add_paragraph('Contents','Heading 1')
                p.paragraph_format.space_before=Pt(0)
                bookmark(p,'Contents',100)
                p=doc.add_paragraph('Select a section to jump to it. Word’s Navigation Pane also lists the topics.')
                p.paragraph_format.space_after=Pt(13)
                for n,h in enumerate(headings,1):
                    p=doc.add_paragraph()
                    p.paragraph_format.space_after=Pt(6)
                    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.65),WD_TAB_ALIGNMENT.RIGHT)
                    hyperlink(p,h,anchor=f'section_{n}')
                    p.add_run('\t');field(p,f'PAGEREF section_{n} \\h','?')
                toc_added=True
            section_count+=1
            p=doc.add_paragraph(line[3:],'Heading 1')
            p.paragraph_format.page_break_before=True
            p.paragraph_format.space_before=Pt(0)
            bookmark(p,f'section_{section_count}',section_count)
            i+=1;continue
        if line.startswith('### '):
            p=doc.add_paragraph(line[4:],'Heading 2')
            if line[4:] == 'Why cost can remain unchanged':
                p.paragraph_format.page_break_before=True
                p.paragraph_format.space_before=Pt(0)
            i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', c) for c in cells): rows.append(cells)
                i+=1
            add_table(doc,rows);continue
        match=re.match(r'^(\d+)\. (.*)',line)
        if match:
            p=doc.add_paragraph()
            p.paragraph_format.left_indent=Inches(.25)
            p.paragraph_format.first_line_indent=Inches(-.25)
            p.add_run(match[1]+'. ').bold=True
            inline(p,match[2]);i+=1;continue
        if line.startswith('- '):
            p=doc.add_paragraph()
            p.paragraph_format.left_indent=Inches(.18)
            p.paragraph_format.first_line_indent=Inches(-.18)
            p.add_run('• ');inline(p,line[2:]);i+=1;continue
        p=doc.add_paragraph();inline(p,line);i+=1
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUTPUT)
    print(f'Created {OUTPUT}')
    print(f'{len(text.split())} words; {len(headings)} sections; {len(doc.tables)} tables')


if __name__=='__main__':build()
